"""Build detailed, editable Korean slides from documented compiler behavior.

No compiler/model runs. slide_content.json owns visible explanations and notes;
this file draws diagrams and renders identical scenes to PPTX and PDF.
"""
from pathlib import Path
import argparse
import hashlib
import json
import math
import sys

HERE=Path(__file__).resolve().parent
DOCS=HERE.parent
SUPPORT=DOCS.parent/'build/presentation_tools'
if (SUPPORT/'deps').exists():sys.path.insert(0,str(SUPPORT/'deps'))
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

W,H=1440,810
C={'bg':'F4F7FB','ink':'172942','muted':'576980','line':'D5DFEC',
   'teal':'008A82','blue':'2869C6','purple':'7655B7','orange':'B46E10','red':'BE5146',
   'sram':'DEF3EE','dram':'FFF0D7','unit':'EBE6F7','ir':'E5EEFB','host':'EBEFF5',
   'white':'FFFFFF','dark':'142941'}
S=[]

def j(path):return json.loads((DOCS/path).read_text())

def wrapped(text,width,size,bold=False):
    font='NPUB' if bold else 'NPU'
    result=[]
    for paragraph in str(text).split('\n'):
        line=''
        for word in paragraph.split(' '):
            candidate=(line+' '+word).strip()
            if pdfmetrics.stringWidth(candidate,font,size)<=width:
                line=candidate;continue
            if line:result.append(line);line=''
            for char in word:
                if pdfmetrics.stringWidth(line+char,font,size)>width:
                    result.append(line);line=''
                line+=char
        result.append(line)
    return result

class Slide:
    def __init__(self,item,index,total):
        self.__dict__.update(item)
        self.ops=[];self.dark=False;self.regions=[];S.append(self)
        self.box(0,0,W,H,'bg',radius=0)
        self.text(48,25,self.section,14,'teal',bold=True,w=1150)
        self.text(1273,26,'상세 설명판',13,'muted',w=120,align='right')
        self.text(48,62,self.title,33,bold=True,w=1344)
        self.text(48,116,self.subtitle or '계산의 의미, 실행 배치, 메모리와 검증을 실제 예제로 연결합니다.',18,'muted',w=1344)
        self.line(48,151,1392,151)
        self.text(48,165,'동작을 따라 읽는 그림',12,'teal',bold=True,w=850)
        self.box(48,191,864,457,'white',radius=12)
        yy=165
        body_size=18
        while sum(49+len(wrapped(t['body'],404,body_size))*body_size*1.38 for t in self.explain)+20>490:
            body_size-=.25
            if body_size<17:raise ValueError(f'Slide {index}: shorten explanation text')
        for i,item in enumerate(self.explain):
            body=wrapped(item['body'],404,body_size)
            height=49+len(body)*body_size*1.38
            self.box(944,yy,448,height,'white',radius=10)
            self.box(944,yy,4,height,['blue','teal','purple'][i],radius=0)
            self.text(964,yy+17,f'{i+1:02d}  '+item['title'],19,bold=True,w=407)
            self.para(965,yy+43,item['body'],404,body_size,color='muted',limit=yy+height-5)
            yy+=height+10
        if yy-10>657:raise ValueError(f'Slide {index}: explanation cards exceed area ({yy})')
        self.box(48,674,1344,80,'sram',radius=10)
        self.text(66,687,'예제와 해석',13,'teal',bold=True,w=124)
        self.para(199,688,self.caption,1169,19,limit=751)
        self.footer(index,total)

    def box(self,x,y,w,h,fill='white',stroke=None,radius=8):
        self.ops.append(dict(kind='box',x=x,y=y,w=w,h=h,fill=C.get(fill,fill),stroke=C.get(stroke,stroke),radius=radius))
    def text(self,x,y,text,size=20,color='ink',bold=False,w=800,align='left'):
        for i,line in enumerate(str(text).split('\n')):
            self.ops.append(dict(kind='text',x=x,y=y+i*size*1.38,w=w,h=size*1.38,text=line,size=size,color=C.get(color,color),bold=bold,align=align))
    def para(self,x,y,text,w,size=18,color='ink',bold=False,limit=None):
        lines=wrapped(text,w,size,bold);height=len(lines)*size*1.38
        if limit is not None and y+height>limit+.5:raise ValueError(f'{self.visual}: text exceeds region: {text}')
        self.text(x,y,'\n'.join(lines),size,color,bold,w)
        self.regions.append(dict(x=x,y=y,w=w,h=height,text=text))
        return height
    def line(self,x1,y1,x2,y2,color='line',width=1.5):
        self.ops.append(dict(kind='line',x1=x1,y1=y1,x2=x2,y2=y2,color=C.get(color,color),width=width))
    def arrow(self,x1,y1,x2,y2,color='muted',width=2):
        self.line(x1,y1,x2,y2,color,width)
        d=math.hypot(x2-x1,y2-y1);ux,uy=(x2-x1)/d,(y2-y1)/d
        self.ops.append(dict(kind='poly',points=[(x2,y2),(x2-ux*8-uy*4,y2-uy*8+ux*4),(x2-ux*8+uy*4,y2-uy*8-ux*4)],fill=C.get(color,color)))
    def route(self,points,color='muted'):
        for a,b in zip(points[:-2],points[1:-1]):self.line(*a,*b,color,2)
        self.arrow(*points[-2],*points[-1],color)
    def card(self,x,y,w,h,title,body='',kind='ir',size=20):
        self.box(x,y,w,h,kind)
        title_lines=wrapped(title,w-28,size,True)
        self.text(x+14,y+14,'\n'.join(title_lines),size,bold=True,w=w-28)
        if body:self.para(x+14,y+20+len(title_lines)*size*1.38,body,w-28,15,color='muted',limit=y+h-8)
    def tag(self,x,y,text,kind='sram',w=185):
        self.box(x,y,w,30,kind,radius=5);self.text(x+9,y+6,text,14,bold=True,w=w-18)
    def table(self,headers,rows,widths,x=67,y=216,row_h=55,size=17):
        self.box(x,y,sum(widths),36,'dark',radius=0)
        for i,(label,width) in enumerate(zip(headers,widths)):
            self.text(x+sum(widths[:i])+10,y+9,label,15,'white',bold=True,w=width-20)
        for r,row in enumerate(rows):
            yy=y+36+r*row_h;self.box(x,yy,sum(widths),row_h,'host' if r%2 else 'bg',radius=0)
            for i,(label,width) in enumerate(zip(row,widths)):
                self.para(x+sum(widths[:i])+10,yy+9,str(label),width-20,size,limit=yy+row_h-3)
    def matrix(self,x,y,w,h,rows,cols,label,highlights=None,kind='ir'):
        self.text(x,y-27,label,18,bold=True,w=w+100)
        for r in range(rows):
            for c in range(cols):
                color=(highlights or {}).get((r,c),kind)
                self.box(x+c*w/cols,y+r*h/rows,w/cols,h/rows,color,'white',0)
        self.box_outline(x,y,w,h)
    def box_outline(self,x,y,w,h,color='line'):
        for a,b in [((x,y),(x+w,y)),((x+w,y),(x+w,y+h)),((x+w,y+h),(x,y+h)),((x,y+h),(x,y))]:self.line(*a,*b,color)
    def bar(self,x,y,width,label,value,maximum,color='teal',display=None):
        self.text(x,y,label,16,w=170)
        self.box(x+172,y,width,26,'host',radius=3)
        self.box(x+172,y,width*value/maximum,26,color,radius=3)
        self.text(x+180+width,y+2,display or f'{value:,}',16,bold=True,w=175)
    def footer(self,index,total):
        self.text(48,781,'NPU compiler · v09 FP16 · 상세 설명'+(' · 부록' if self.appendix else ''),10,'muted',w=360)
        labels=[p.split('/')[-1] for p in self.sources]
        self.text(412,781,'근거: '+', '.join(labels),9,'muted',w=895)
        self.text(1325,779,f'{index:02d} / {total:02d}',12,'muted',w=65,align='right')

def flow(s,items,y=230,x=70,w=246,h=104,gap=43):
    for i,(title,body,kind) in enumerate(items):
        xx=x+i*(w+gap);s.card(xx,y,w,h,title,body,kind)
        if i:s.arrow(xx-gap+4,y+h/2,xx-7,y+h/2)

def code(s,x,y,w,h,title,lines,kind='host',size=17):
    s.box(x,y,w,h,kind);s.text(x+17,y+13,title,18,'teal',bold=True,w=w-34)
    s.text(x+17,y+49,lines,size,w=w-34)

def draw(s):
    v=s.visual
    if v=='1':
        s.text(77,219,'한 모델이 실행되는 세 가지 관점',26,bold=True,w=780)
        flow(s,[('모델의 계산','X → attention → MLP\n28 layers → logits','ir'),('실행 배치','64×64 tile + panel\nFP32 누산 → FP16 저장','unit'),('값의 보관','SRAM 재사용 / DRAM\nKV cache와 주소 계약','sram')],y=283,w=247,gap=42,h=130)
        s.route([(194,413),(194,451),(772,451),(772,482)],'teal')
        s.card(70,482,821,126,'IR와 명령, 실험 결과로 선택을 확인','작은 예제: 실제 shape·수명·주소를 읽기\nLlama: hidden 3072 / MLP 8192 / Q 24 heads / KV 8 heads',kind='host',size=23)
    elif v=='2':
        s.table(['단계','같은 예제가 어떻게 보이는가','이번 단계에서 정하는 것'],[
            ['계산 의미','Y = SiLU(X × B + bias)','연산 연결 · reduction · cast'],
            ['실행 배치','출력 64×64 / K를 64씩 2번','Loop 순서 · cache 위치'],
            ['저장과 주소','입력 X/B 각 16 KiB를 SRAM에','수명 · 재사용 · 실제 offset'],
            ['실행과 검증','M_MUL + vector + DMA stream','ABI · 출력 비교 · 비용']], [148,357,321],row_h=77,size=18)
        s.tag(75,581,'같은 계산을 단계별로 구체화','sram',360)
    elif v=='3':
        flow(s,[('COMPILE','정적 shape / dtype\n장치 용량과 수치 계약','ir'),('PREPARE','Checkpoint 배열\nHost transform 실행','host'),('RUN','입력 patch / 실행\nLogits 읽기 / argmax','sram')],y=231,w=247,gap=42,h=140)
        for x,a,b,k in [(70,'Program + plan','명령과 주소의 연결','ir'),(359,'DRAM image','변환된 W + 입력 + 상수','dram'),(648,'다음 호출의 상태','선택 token + 갱신된 KV','sram')]:
            s.arrow(x+123,371,x+123,414);s.card(x,414,247,104,a,b,k,19)
        s.route([(775,518),(775,580),(660,580),(660,530)],'teal')
        s.text(86,579,'Decode의 기본 계산 명령은 step마다 재생성하지 않음',19,'teal',bold=True,w=550)
    elif v=='4':
        s.card(71,216,560,105,'DRAM: 큰 tensor와 persistent 상태','W 전체 48 MiB · 외부 입출력 · KV cache','dram',23)
        s.card(71,366,560,106,'SRAM: 기본 8 MiB','입력 panel | 다음 kernel에 남길 값 | scratch','sram',23)
        s.card(71,517,267,105,'Matrix unit','64×64 / FP32 accumulator','unit',22)
        s.card(360,517,271,105,'Vector unit','Pointwise + reduction','unit',22)
        s.arrow(245,321,245,366,'orange');s.arrow(435,366,435,321,'orange')
        s.text(652,323,'GLOAD / GSTORE',21,'orange',bold=True,w=235)
        s.arrow(195,472,195,517,'purple');s.arrow(490,517,490,472,'purple')
        s.text(652,480,'LOAD / SAVE',21,'purple',bold=True,w=235)
        s.para(653,537,'SRAM 재사용은 DRAM 이동을 줄임. Unit LOAD는 여전히 필요.',219,18)
    elif v=='5':
        s.tag(70,208,'입력 X [S,3072]','sram',211)
        s.card(72,266,143,88,'RMSNorm','마지막 축 정규화','ir',19)
        s.card(250,266,241,88,'Q / K / V projection','Q:[S,24,128]  KV:[S,8,128]','ir',18)
        s.card(527,266,167,88,'Attention','RoPE / QK / PV','unit',16)
        s.card(729,266,155,88,'O projection','[S,3072]','ir',19)
        for a,b in [(215,250),(491,527),(694,729)]:s.arrow(a,310,b,310)
        s.card(660,390,224,72,'R = X + attention','','sram',19)
        s.route([(805,354),(805,390)])
        s.route([(166,238),(166,249),(58,249),(58,426),(660,426)],'teal')
        s.arrow(166,249,166,266,'teal')
        s.text(92,380,'Residual X를 add까지 유지',17,'teal',w=480)
        s.card(70,502,157,86,'RMSNorm(R)','[S,3072]','ir',18)
        s.card(273,483,260,61,'Gate → SiLU  [S,8192]','','unit',18)
        s.card(273,559,260,61,'Up  [S,8192]','','ir',18)
        s.route([(660,430),(591,430),(591,473),(149,473),(149,502)])
        s.route([(227,545),(249,545),(249,513),(273,513)])
        s.route([(249,545),(249,589),(273,589)])
        s.card(576,521,138,63,'× → Down','','unit',17)
        s.route([(533,513),(555,513),(555,550),(576,550)])
        s.route([(533,589),(555,589),(555,552),(576,552)])
        s.card(755,521,129,63,'Y = R + MLP','','sram',16)
        s.arrow(714,552,755,552)
        s.route([(884,426),(897,426),(897,551),(884,551)],'teal')
    elif v=='6':
        s.table(['단계','TVM 표준 기반','NPU가 추가하는 결정'],[
            ['그래프','Bindings / CSE / LiftTransformParams','Attention 의미·지원 layout'],
            ['계산 정의','LegalizeOps / TE','FP32 누산과 FP16 저장'],
            ['Fusion','Annotate / FuseOps / FuseTIR','지원 영역 / 부분·전체 비용'],
            ['Schedule','split / reorder / cache / tensorize','Panel·복사 위치·장치 intrinsic'],
            ['Memory','StaticPlanBlockMemory / USMP','호출 간 보관 + schedule 재선택'],
            ['실행','IR 분석 기반 codegen 연결','ISA / linker / runtime ABI']], [132,357,337],y=211,row_h=59,size=17)
    elif v=='7':
        s.matrix(78,260,155,110,2,4,'X [64,128]',kind='sram')
        s.text(249,303,'×',28,w=35)
        s.matrix(299,260,115,155,4,2,'B [128,64]',kind='ir')
        s.arrow(429,314,466,314)
        s.matrix(484,260,137,110,2,2,'P [64,64]',kind='unit')
        s.text(645,301,'+',28,w=35)
        s.matrix(705,294,166,39,1,4,'bias [64]',kind='dram')
        s.route([(772,350),(772,437),(517,437),(517,465)])
        s.route([(554,370),(554,437),(517,437),(517,465)])
        flow(s,[('A = P + bias','[64,64] FP16 / 8 KiB','sram'),('Y = SiLU(A)','[64,64] FP16 / 8 KiB','unit')],y=466,x=389,w=234,gap=36,h=100)
        s.card(77,464,270,135,'B = W.T','W:[64,128] → B:[128,64]\nParameter-only host 처리','host',20)
    elif v=='8':
        code(s,72,224,368,166,'변환 전: 매 호출 graph','device(x, W, bias)\n  B = transpose(W)\n  y = SiLU(x @ B + bias)',size=18)
        s.arrow(453,304,495,304,'teal')
        code(s,512,223,371,166,'LiftTransformParams 이후','host_transform(W, bias)\n  return (bias, transpose(W))','host',18)
        code(s,512,420,371,157,'Device 입력 signature','device(x, bias, B)\n  y = SiLU(x @ B + bias)','ir',18)
        s.arrow(697,389,697,420,'teal')
        s.card(73,423,368,152,'실제 예제에서 달라진 순서','원래: x / W / bias\n변환 후: x / bias / lv (= B)\nRuntime이 새 signature에 맞춰 연결','sram',20)
    elif v=='9':
        s.card(73,221,810,83,'Relax attention(Q,K,V, bias, scale, mask)','Layout과 causal/window 의미를 보존하여 세부 연산으로 변환','ir',22)
        s.arrow(472,305,472,346,'teal')
        for i,(title,detail,kind) in enumerate([('QK^T','scores','unit'),('Scale','1 / sqrt(d)','ir'),('Mask','bias / causal','ir'),('Softmax','행별 확률','unit'),('PV','head 출력','unit')]):
            x=73+i*165;s.card(x,347,148,98,title,detail,kind,20)
            if i:s.arrow(x-17,396,x-3,396)
        s.box_outline(221,334,482,127,'blue');s.text(260,474,'Score fusion 후보: scale + mask + softmax',18,'blue',bold=True,w=580)
        s.line(77,516,879,516,'teal',3);s.text(215,531,'Full 후보: QK^T부터 PV까지 함께 schedule',20,'teal',bold=True,w=650)
        s.text(90,593,'분해 = 수식 공개    /    Fusion = 함께 실행할 영역 선택',18,'muted',w=786)
    elif v=='10':
        code(s,73,221,811,155,'하나의 IRModule 안에 함께 존재하는 Relax 함수','p = R.call_tir(matmul, (x, B), out_sinfo=...)\na = R.call_tir(add, (p, bias), out_sinfo=...)\ny = R.call_tir(silu, (a,), out_sinfo=...)','ir',18)
        for i,(title,body) in enumerate([('matmul PrimFunc','i,j: 출력 위치\nk: reduction 축\nFP32 누산 → FP16 cast'),('add PrimFunc','각 i,j의 p를 읽음\nbias[j] broadcast\nFP16 결과 저장'),('silu PrimFunc','각 원소에 수식 적용\nx / (1 + exp(-x))\nFP16 중간값')]):
            x=73+i*278;s.arrow(x+126,376,x+126,416);s.card(x,416,254,179,title,body,'unit',21)
    elif v=='11':
        s.matrix(77,257,213,103,2,4,'X: K=128',highlights={(r,c):'sram' if c<2 else 'ir' for r in range(2) for c in range(4)})
        s.matrix(347,257,103,162,4,2,'B: K=128',highlights={(r,c):'sram' if r<2 else 'ir' for r in range(4) for c in range(2)})
        s.text(95,377,'K 0..63 | 64..127',17,w=240)
        s.card(503,219,378,91,'1. 첫 64개: MAC=0','FP32 accumulator를 새로 시작','sram',21)
        s.card(503,335,378,91,'2. 다음 64개: MAC=1','같은 출력에 FP32 부분합 추가','ir',21)
        s.card(503,452,378,95,'3. K 완료: FP16 SAVE','새 출력 tile이면 다음 명령은 MAC=0','unit',21)
        s.arrow(691,310,691,335);s.arrow(691,426,691,452)
        s.tag(79,499,'K 중간 FP16 SAVE는 하지 않음','dram',383)
        s.text(83,580,'현재 자동 matmul: activation OFF / 뒤의 SiLU는 vector 경로',19,'purple',bold=True,w=801)
    elif v=='12':
        for y,label,groups in [(221,'A · 세 kernel',[(0,1),(1,1),(2,1)]),(345,'B · 두 kernel',[(0,1),(1,2)]),(469,'C · 한 kernel',[(0,3)])]:
            s.text(73,y,label,20,bold=True,w=240)
            for start,count in groups:s.box(309+start*191,y-7,count*191-15,83,'sram' if count>1 else 'ir',stroke='teal' if count>1 else None)
            for i,t in enumerate(['matmul','add','SiLU']):
                s.text(327+i*191,y+19,t,22,bold=True,w=150,align='center')
                if i<2:s.arrow(476+i*191,y+36,493+i*191,y+36)
        s.tag(81,584,'A/B도 SRAM 중간 전달 가능','sram',356)
        s.tag(462,584,'많이 합침 ≠ 항상 낮은 비용','dram',395)
    elif v=='12a':
        s.card(73,216,220,84,'Matmul','pattern 4','unit',21)
        s.card(368,216,211,84,'Add / SiLU','pattern 1 / 0','ir',21)
        s.card(660,216,222,84,'후속 소비','의존 관계 검사','sram',21)
        s.arrow(293,258,368,258);s.arrow(579,258,660,258)
        s.route([(321,258),(321,343),(476,343)],'orange')
        s.card(478,317,404,84,'외부 출력 또는 다른 사용자','값을 밖에서도 관찰하면 그 사용을 보존','dram',20)
        s.table(['분류','예','Fusion에서 확인할 것'],[
            ['0 / 1 / 2','원소별 / broadcast / injective','연결 경로와 같은 group의 제약'],
            ['3 / 4','Reduction / matmul 계열','임의의 복합 연산끼리 합치지 않음'],
            ['7 / 8','Tuple / opaque','Tuple 관계 / 일반 fusion 경계']], [159,260,407],y=437,row_h=49,size=16)
    elif v=='13':
        flow(s,[('영역 하나 검사','TVM이 만든 작은 범위\nfused ↔ unfused 비교','ir'),('전체 후보 조립','살아남은 영역 + 주변 graph\nprefill / decode entry','host'),('전체 module 선택','보관 + 실제 주소 + 최종 명령\n모든 entry 비용 확인','sram')],y=226,w=247,gap=42,h=141)
        for x,label in [(72,'부분 검사도 codegen까지'),(650,'전체 수명과 ABI까지')]:s.tag(x,390,label,'unit',238)
        s.card(74,445,807,158,'비용 비교의 우선순위','1. DMA bytes가 작으면 우선\n2. DMA가 같으면 ISA words, 다시 같으면 SRAM peak\nPrefill/decode 중 어떤 entry도 나빠지지 않는 개선을 채택','sram',21)
    elif v=='14':
        s.matrix(86,260,297,157,2,4,'X panel [64,128] · 16 KiB',highlights={(r,c):'sram' if c<2 else 'ir' for r in range(2) for c in range(4)})
        s.matrix(494,260,157,230,4,2,'B panel [128,64] · 16 KiB',highlights={(r,c):'sram' if r<2 else 'ir' for r in range(4) for c in range(2)})
        s.text(410,318,'×',36,w=46)
        s.card(697,277,184,150,'Unit tile','64×64\nK 64개씩\n두 번 누산','unit',21)
        s.arrow(656,350,690,350,'purple')
        s.tag(80,456,'GLOAD: 두 panel을 한 번씩','dram',333)
        s.card(76,532,810,77,'K loop 안에서는 SRAM의 다른 부분을 LOAD','DRAM panel 복사 위치와 연산 tile의 반복 위치를 구분','sram',20)
    elif v=='15':
        code(s,72,225,370,225,'변환 전: B 전체를 먼저 생성','for i in 0..3:\n  B[i,:] = A[i,:] + 1\n\nfor i in 0..3:\n  C[i,:] = B[i,:] * 2',size=19)
        code(s,508,225,376,225,'compute_at 뒤: 필요한 두 행','for io in [0, 2]:\n  B[io:io+2,:] = A[...] + 1\n  C[io:io+2,:] = B[...] * 2',kind='sram',size=18)
        s.arrow(451,333,499,333,'teal')
        for x in [94,529]:
            s.matrix(x,526,259,74,4,4,'B가 필요한 범위',highlights={(r,c):'sram' if x==94 or r<2 else 'host' for r in range(4) for c in range(4)})
        s.text(813,537,'2행씩',17,'teal',bold=True,w=69)
    elif v=='15a':
        for y,title,a,b,kind in [(223,'compute_at(B, C의 loop)','C의 소비 loop가 기준','B 생산을 loop 안으로','ir'),(431,'reverse_compute_at(C, B의 loop)','B의 생산 loop가 기준','C 소비를 loop 안으로','sram')]:
            s.text(77,y,title,22,'teal',bold=True,w=800)
            s.card(78,y+49,295,123,a,'필요한 행/타일의 범위\n의존성을 유지하며 배치',kind,20)
            s.arrow(389,y+110,443,y+110)
            s.card(462,y+49,416,123,b,'for tile:\n    B tile 생성 → C tile 소비',kind,20)
    elif v=='15b':
        flow(s,[('cache_read','입력 → 새 cache buffer\n소비자는 cache를 읽음','ir'),('compute_at','복사를 행 묶음 안으로\n매번 필요한 범위만 접근','sram'),('Compaction','실제 allocation 축소\n4×4 → 2×4 예제','unit')],y=222,w=247,gap=42,h=141)
        s.matrix(91,432,165,92,4,4,'처음 cache',kind='ir')
        s.arrow(281,472,368,472,'teal')
        s.matrix(402,432,165,92,4,4,'접근 범위',highlights={(r,c):'sram' if r<2 else 'host' for r in range(4) for c in range(4)})
        s.arrow(592,472,679,472,'teal')
        s.matrix(715,432,165,46,2,4,'축소 뒤',kind='sram')
        s.tag(78,576,'set_scope = 기존 buffer의 scope','host',390)
        s.tag(489,576,'cache_write = 출력 cache + 복사','host',391)
    elif v=='16':
        s.table(['선택 축','후보의 예 / 뜻'],[
            ['m_panel / n_panel','출력 행·열 묶음: N=256이면 64 / 128 / 256'],
            ['k_chunk','K를 묶는 단위; 안에서는 64개씩 matrix 실행'],
            ['outer / inner order','M/N panel 순서와 panel 안 M/N tile 순서'],
            ['a_at / b_at','X/W 복사를 둘 loop 위치를 각각 선택']], [239,587],y=216,row_h=53,size=17)
        for i,(t,k) in enumerate([('추정 DMA 순위','ir'),('Schedule 생성','unit'),('공간·누산 검사','dram'),('최종 비용 비교','sram')]):
            x=73+i*207;s.card(x,494,190,65,t,'',k,18)
            if i:s.arrow(x-17,526,x-3,526)
        s.text(79,589,'보관값이 차지한 공간을 제외 → 남은 SRAM으로 다시 후보 평가',19,'teal',bold=True,w=791)
    elif v=='17':
        s.matrix(80,263,208,150,2,6,'X [64,3072] 384 KiB',kind='sram')
        s.matrix(370,263,504,150,3,8,'W [3072,8192] 전체 48 MiB',highlights={(r,c):'sram' if c==0 else 'host' for r in range(3) for c in range(8)})
        s.text(81,436,'X 유지',20,'teal',bold=True,w=240)
        s.text(373,436,'64열 panel 384 KiB씩 교체 × 128회',20,'teal',bold=True,w=504)
        s.card(76,494,392,117,'Gate / up의 matrix 명령','출력 tile 128개 × K tile 48개\n= 6144개 M_MUL','unit',21)
        s.card(491,494,392,117,'Down의 기록된 선택','W 4 MiB panel에 출력 tile 4개\n12 panels × 4 × 128 = 6144','sram',21)
    elif v=='18':
        s.card(73,219,162,68,'입력 x','','sram',21)
        stages=[('-x','부호 반전'),('exp','지수'),('1 + t','분모'),('1 / t','역수'),('x × t','SiLU 출력')]
        for i,(t,b) in enumerate(stages):
            x=73+i*165;s.card(x,348,148,105,t,b,'unit' if i%2 else 'ir',22)
            if i:s.arrow(x-17,400,x-4,400)
        s.route([(150,287),(150,348)],'teal')
        s.route([(235,253),(806,253),(806,348)],'teal')
        s.text(398,277,'마지막 곱셈까지 x가 필요',19,'teal',bold=True,w=429)
        s.card(76,512,805,99,'각 단계: SRAM에서 LOAD → vector 계산 → SRAM에 SAVE','임시값 수명·길이에 맞춰 scratch를 재사용하고 연속 축은 길게 처리','sram',21)
    elif v=='18a':
        s.text(78,220,'논리 VLEN = 1000개 원소',25,bold=True,w=800)
        widths=[206,206,206,186];xx=76
        for i,(n,w) in enumerate(zip([256,256,256,232],widths)):
            s.card(xx,282,w,97,str(n)+'개','물리 폭 단위' if i<3 else '유효 tail','ir' if i<3 else 'sram',25);xx+=w+8
        s.text(89,402,'0..255',17,w=190);s.text(301,402,'256..511',17,w=195);s.text(515,402,'512..767',17,w=195);s.text(729,402,'768..999',17,w=169)
        s.card(77,466,803,145,'현재 계약의 전제와 아직 확인할 것','Compiler/C-model: 긴 논리 길이 명령을 표현\nRTL 확인 필요: 256-lane 내부 반복, tail, LOAD/SAVE와 reduction 누적\n명령 1개라도 처리 길이에 따라 실제 실행 cycle은 달라질 수 있음','dram',21)
    elif v=='18b':
        rows=[['입력','[2, 1, 0]','FP16'],['Max와 빼기','max=2 → [0, -1, -2]','행별 reduction + broadcast'],['Exp','[1, 0.3679, 0.1353]','FP16 vector 결과'],['Sum','약 1.5032','FP32 합산 → FP16 저장'],['Normalize','[0.6652, 0.2447, 0.0900]','각 exp / 합계']]
        s.table(['계산 단계','설명용 값','연산 / 정밀도'],rows,[164,368,294],y=222,row_h=67,size=18)
    elif v=='18c':
        s.card(76,219,803,79,'문제: 1000² = 1,000,000 → FP16 저장 범위를 초과','제곱을 FP16으로 만든 뒤 FP32로 더하는 것만으로는 해결되지 않음','dram',21)
        flow(s,[('1. 크기 축소','a = max(absmax(x), √eps)\nu = x / a  → |u| ≤ 1','sram'),('2. 안전한 제곱·합','z = u / √width\ns = sum(z × z)','ir'),('3. 보정하여 출력','inv = 1/√(s+(√eps/a)²)\ny = u × inv × weight','unit')],y=347,w=247,gap=42,h=135)
        s.card(76,535,803,78,'수식은 실수에서 등가, 실제 FP16 계산은 target 근사','LegalizeOps에서 정의한 수치 정책이며 자동 정밀도 탐색은 아님','host',21)
    elif v=='18d':
        s.card(77,224,371,143,'RoPE: head의 앞/뒤 절반','a = Q 또는 K의 앞 64개\nb = 뒤 64개\nV는 회전하지 않음','ir',21)
        s.card(488,224,392,143,'같은 position의 cos / sin','a′ = a × cos - b × sin\nb′ = b × cos + a × sin\nHost가 위치별 cos/sin 준비','unit',21)
        s.arrow(451,294,485,294)
        s.text(79,415,'GQA: KV head 하나를 Q heads 세 개가 공유',23,bold=True,w=800)
        for i in range(3):s.card(79+i*130,470,110,63,f'Q {i}','','ir',21)
        s.card(570,468,300,77,'KV head 0','24 Q / 8 KV = 3개씩 공유','sram',20)
        for i in range(3):s.route([(132+i*130,535),(132+i*130,580+i*14),(720,580+i*14),(720,545)],'teal')
    elif v=='19':
        s.matrix(82,256,340,252,8,10,'Score [S,S] · head 하나',highlights={(r,c):'sram' if r==2 else 'host' for r in range(8) for c in range(10)})
        s.text(100,533,'열 방향: 모든 key 길이 S',19,bold=True,w=341)
        s.text(89,574,'선택한 query 64행만 SRAM에 생성',18,'teal',bold=True,w=408)
        s.arrow(426,335,478,335,'teal')
        s.card(493,221,385,102,'Q 조각 × K^T → score 조각','S=3072: 64×3072×2 = 384 KiB','sram',21)
        s.card(493,352,385,98,'Scale / mask / softmax','한 query 행의 key 전체를 처리','unit',21)
        s.card(493,479,385,99,'P 조각 × V → 출력 조각','다음 query 묶음으로 이동','ir',21)
        s.arrow(687,323,687,352);s.arrow(687,450,687,479)
        s.text(513,601,'Key online softmax는 현재 없음',18,'red',bold=True,w=366)
    elif v=='20':
        s.table(['범위','분석과 출력','어떤 문제를 푸는가'],[
            ['Relax 전체 함수','StaticPlanBlockMemory\n→ storage pool','수명이 겹치지 않는 tensor의\n기본 DRAM 공간 재사용'],
            ['Scheduled TIR\n한 kernel 안','USMP: allocation 충돌\n→ byte offset 배치','임시 buffer끼리 공간 공유\n장치 상태는 보수적으로 보호'],
            ['Kernel 사이\n호출 순서','NPU memory search\n→ 보관 / DRAM 선택','중간값을 마지막 소비까지 보관\n남은 용량으로 schedule 재선택']], [176,294,356],y=220,row_h=103,size=18)
        s.tag(79,592,'서로 다른 범위의 세 결정이 합쳐져 최종 주소가 됨','sram',791)
    elif v=='21':
        x0=252;step=123
        for i,t in enumerate(['gate','SiLU','up','multiply','down']):
            s.text(x0+i*step,217,t,17,bold=True,w=117,align='center');s.line(x0+i*step,254,x0+i*step,568)
        for r,(name,a,b,kind) in enumerate([('G · 1 MiB',0,1,'ir'),('S · 1 MiB',1,3,'sram'),('U · 1 MiB',2,3,'unit'),('H · 1 MiB',3,4,'dram')]):
            y=279+r*68;s.text(79,y+9,name,21,bold=True,w=156)
            s.box(x0+a*step+6,y,(b-a)*step+105,44,kind)
            s.text(x0+a*step+15,y+13,'생성 → 마지막 소비',15,w=(b-a)*step+85)
        s.text(262,599,'시간 / 실행 순서 →',18,'teal',bold=True,w=351)
        s.tag(640,591,'겹친 수명은 같은 주소 불가','dram',241)
    elif v=='22':
        for y,title,parts in [(253,'Kernel 1: X × W1 → H',[('상수 등',1548,'host'),('X',8192,'ir'),('W1',16384,'unit'),('H',16384,'sram')]),(467,'Kernel 2: H × W2 → Y',[('상수 등',1548,'host'),('Y 작업',8192,'ir'),('W2',16384,'unit'),('같은 H',16384,'sram')])]:
            s.text(80,y-43,title,23,bold=True,w=789)
            xx=79
            # Constant strip exaggerated for visibility; other segments follow relative sizes.
            for label,nb,kind in parts:
                ww=80 if nb==1548 else (718*nb/40960)
                s.box(xx,y,ww,105,kind,'white',0)
                s.text(xx+5,y+20,label,16,bold=True,w=ww-10,align='center')
                s.text(xx+5,y+56,str(nb)+' B',14,w=ww-10,align='center');xx+=ww
        for xx,label in [(79,'0'),(159,'1548'),(302.6,'9740'),(589.8,'26124'),(877,'42508')]:s.text(xx-8,376,label,14,'muted',w=78)
        s.arrow(874,359,874,464,'teal',4)
        s.text(322,411,'공간 재사용',18,'purple',bold=True,w=203)
        s.text(606,412,'H 보관 / DRAM 왕복 없음',18,'teal',bold=True,w=267)
        s.text(84,604,'상수 영역은 식별을 위해 확대 표시 / 모든 주소·크기는 bytes',16,'muted',w=794)
    elif v=='23':
        s.text(80,220,'H [64,128] FP16: 한 행 = 256 bytes',23,bold=True,w=791)
        s.matrix(81,280,390,186,4,2,'',highlights={(r,c):'sram' if c==0 else 'ir' for r in range(4) for c in range(2)})
        s.text(104,480,'왼쪽 64열',20,'teal',bold=True,w=168)
        s.text(299,480,'오른쪽 64열',20,'blue',bold=True,w=199)
        s.text(91,525,'시작: H',18,w=180);s.text(288,525,'시작: H + 128 B',18,w=224)
        s.card(518,272,358,106,'완료된 accumulator를 직접 SAVE','부분 영역 view가 전체 H의\nrow stride 256 B를 유지','unit',19)
        s.card(518,411,358,148,'DRAM 예약 정리','초기: H용 DRAM pool 존재\n선택: H를 SRAM-only로 전달\n후처리: 완전히 미사용인 pool 제거','dram',20)
        s.tag(89,584,'Tile 8 KiB를 단순 연속 배치한 주소와 다름','sram',775)
    elif v=='24':
        flow(s,[('Semantic TIR','C32 += fp32(X) × fp32(B)\nY = fp16(C32)','ir'),('Scheduled TIR','64×64 block tensorize\nnpu_gemm_acc(...)','sram'),('ISA stream','Descriptor / LOAD\nM_MUL / SAVE','unit')],y=229,w=247,gap=42,h=137)
        s.table(['실제 word offset','명령의 의미','누산 관계'],[
            ['30 / 44','첫 matmul의 H 왼쪽 / 오른쪽 tile','둘 다 MAC=0: 다른 출력'],
            ['60 / 67','두 번째 matmul의 K tile 두 개','MAC=0 다음 MAC=1'],
            ['71 / 72 / 76','Y SAVE / GSTORE / HALT','K 완료 뒤 외부 출력 저장']], [177,381,268],y=416,row_h=57,size=17)
    elif v=='24a':
        s.text(80,221,'X GLOAD · word offset 8 · 4-word packet',23,bold=True,w=807)
        for i,t in enumerate(['000c18a0','00000000','00000800','00010800']):s.card(77+i*205,274,190,63,t,'','ir',23)
        s.table(['필드 / 위치','원래 값과 단위','Byte로 읽기'],[
            ['SRAM 주소','0xC18 = 3096 nibbles','3096 ÷ 2 = 1548 B'],
            ['DRAM 시작','0 cells','0 × 4 = 0 B'],
            ['전송량','0x800 = 2048 cells','2048 × 4 = 8192 B'],
            ['다른 필드 주의','Plan은 기본 2 B / stride는 원소','FP16 stride 128 = 256 B']], [213,299,314],y=377,row_h=52,size=17)
    elif v=='25':
        s.card(77,229,365,122,'program_memory.bin','계산 순서와 주소가 인코딩된\nlittle-endian uint32 명령 stream','unit',23)
        s.card(514,229,365,122,'global_memory.bin','Plan에 맞춰 입력·parameter·상수를\n넣은 초기 DRAM image','dram',23)
        s.card(77,430,365,153,'IR dump','lowered.py: 선택된 계산\nscheduled_*.py: 실제 tile/cache 배치\n결정이 코드에 어떻게 반영됐는지 확인','ir',23)
        s.card(514,430,365,153,'program.json','I/O 이름·shape·dtype·offset\n비용 / 선택 사유 / source hash\nRuntime이 결과를 읽는 계약','sram',23)
        s.route([(259,351),(259,390),(697,390),(697,430)],'teal')
        s.arrow(697,351,697,430,'teal')
    elif v=='26':
        s.table(['호출','NPU에 넣은 token / position','Cache 유효 길이','예측 token'],[
            ['Prefill','Prompt 7 tokens','7','358 · “ I”'],
            ['Decode 1','358 / position 7','8','2846 · “\'m”'],
            ['Decode 2','2846 / position 8','9','4560 · “ trying”']], [139,326,162,199],y=223,row_h=76,size=18)
        s.card(80,518,800,93,'Host 준비 → NPU 28 layers → host argmax → 다음 입력','마지막 예측 token 4560은 아직 입력으로 처리되지 않았으므로 그 K/V도 아직 없음','sram',21)
    elif v=='27':
        s.text(78,218,'이번 attention의 입력: capacity 9 + 현재 token 1',22,bold=True,w=811)
        for i in range(10):
            x=78+i*81;k='sram' if i<7 else ('unit' if i==9 else 'host')
            s.card(x,281,73,89,str(i),'과거' if i<7 else '현재' if i==9 else '빈칸',k,23)
            s.text(x+8,393,'유효' if i<7 or i==9 else '제외',15,'teal' if i<7 or i==9 else 'red',w=61)
        s.card(78,487,800,99,'Persistent cache에 기록: 첫 빈 슬롯 L=7','Attention에서 현재 token의 임시 위치는 C=9, 보존할 최종 위치는 L=7','dram',21)
        s.route([(845,374),(845,446),(686,446),(686,487)],'purple')
        s.text(79,607,'Mask는 7,8번의 기여를 막음 / 고정 shape의 연산량이 자동으로 줄지는 않음',17,'muted',w=809)
    elif v=='28':
        flow(s,[('새 K/V 계산 결과','별도 DRAM 출력\nK/V 합계 114688 B','unit'),('NPU DMA append','Host가 position 주소 patch\n기록된 DMA 458752 B','sram'),('Persistent cache','Layer·head별 슬롯 L\n다음 decode 호출에 사용','dram')],y=230,w=247,gap=42,h=153)
        s.table(['층','기록된 동작','섞어 세면 안 되는 이유'],[
            ['Host patch','작은 입력 6676 B / suffix 갱신','Weight rewrite 절감과 관련'],
            ['NPU DMA','위치·layout에 따른 cache 이동','실제 명령의 DRAM 전송'],
            ['Simulator I/O','매 호출 전체 image 파일 read/write','C-model transport 비용']], [145,355,326],y=435,row_h=55,size=17)
    elif v=='29':
        s.card(74,219,805,105,'구조·profile·메모리 조건이 같은가?','Structural hash + equality + 조건 확인 → schedule / 영역 / entry 비용 cache','ir',22)
        for i in range(4):s.card(76+i*205,380,184,100,'Layer '+str(i+1) if i<3 else '… Layer 28','다른 weight 주소\n모든 계산 실행','sram',22)
        s.line(175,341,790,341,'teal',2)
        for x in [168,374,579,790]:s.arrow(x,341,x,380,'teal')
        s.tag(81,543,'같은 판단 재사용 ≠ 실행 layer 생략','unit',785)
        s.text(87,597,'첫 합법 모듈 완성 → 기본 120초 동안 추가 graph 후보 탐색',19,'blue',bold=True,w=798)
    elif v=='30':
        s.table(['검증 / 지표','확인하는 것','해석의 한계'],[
            ['Bit equality / cosine','수치 의미 / 모델 결과','짧은 입력이 모든 품질을 대표하지 않음'],
            ['DMA bytes','명령이 요청한 DRAM 이동','Host patch / 파일 I/O와 다름'],
            ['ISA words','명령 저장 공간','Hardware latency와 다름'],
            ['SRAM peak / image','최고 SRAM / DRAM 예약 크기','전송량과 다름'],
            ['Compile seconds','Compiler의 wall-clock 비용','Inference 속도와 다름']], [220,260,346],y=225,row_h=65,size=18)
    elif v=='31':
        r=j('examples/pipeline/stage_summary.json')['records'];a=r['off']['statistics'];b=r['auto']['statistics']
        s.text(80,220,'ISA words · 명령 저장 크기',22,bold=True,w=800)
        s.bar(80,268,380,'Fusion off / 3 kernels',a['words'],3000,'blue')
        s.bar(80,312,380,'Fusion auto / 1 kernel',b['words'],3000,'teal')
        s.table(['지표','Off','Auto','결과'],[
            ['DMA bytes',a['dma_bytes'],b['dma_bytes'],'같음'],
            ['SRAM peak',a['sram_peak_bytes'],b['sram_peak_bytes'],'8192 B 감소'],
            ['출력','FP16','FP16','Bit-exact / LLVM max diff 0']], [206,156,156,308],y=389,row_h=59,size=18)
        s.tag(81,602,'같은 전송량에서도 명령과 작업 공간은 줄어들 수 있음','sram',792)
    elif v=='32':
        rows=j('examples/memory/measurements.json')['records']
        s.text(78,222,'DMA bytes · 두 matmul / fusion off',23,bold=True,w=811)
        for i,r in enumerate(rows):
            label=f"{r['sram_capacity_bytes']//1024} KiB / {'보관 on' if r['joint_memory'] else '보관 off'}"
            s.bar(80,275+i*58,410,label,r['statistics']['dma_bytes'],90000,'teal' if r['statistics']['dma_bytes']<80000 else 'blue')
        s.tag(83,593,'H 16 KiB의 저장 + 재읽기 = 32768 B 제거','sram',788)
    elif v=='33':
        s.table(['기록 / 범위','Words 전 → 후','DRAM image 전 → 후'],[
            ['E3 · 두 matmul\n실행 / bit-exact','85 → 77','65548 → 49164 B\n16 KiB 예약 제거'],
            ['E4 · 실제 폭 MLP\n정적 compile','143404 → 140594','154927116 → 151781388 B\n3 MiB 예약 제거']], [238,245,343],y=228,row_h=108,size=19)
        s.card(77,521,803,94,'두 비교 모두 DMA와 SRAM peak는 그대로','이미 SRAM으로 전달하던 중간값의 불필요한 명령·DRAM 예약을 정리한 결과','sram',21)
    elif v=='34':
        rows={r['label']:r for r in j('evidence/current/compile_comparison.json')['rows']}
        for y,t,keys in [(223,'M=7 · 실패 검사와 중복 시도 감소',['instrumented_mlp7','bounded_mlp7']),(432,'M=64 · 두 경우 모두 47 trials',['unbounded_mlp64','bounded_mlp64'])]:
            s.text(80,y,t,22,bold=True,w=800)
            for i,k in enumerate(keys):s.bar(82,y+51+i*47,390,'이전' if i==0 else '이후',rows[k]['seconds'],90,'blue' if i==0 else 'teal',f"{rows[k]['seconds']:.3f} s")
        s.tag(80,600,'두 비교 모두 최종 binary hash 동일 / 과거 단일 관측','sram',799)
    elif v=='35':
        s.table(['실행','HF와 같은 token','Cosine','최대 절대차'],[
            ['Prefill','358','0.9999977148','0.03515625'],
            ['Decode pos 7','2846','0.9999958473','0.02343750'],
            ['Decode pos 8','4560','0.9999929546','0.03906250']], [209,194,228,195],y=230,row_h=70,size=19)
        s.card(80,531,797,84,'생성: “ I\'m trying”','유한 logits + token 일치 + cosine ≥ 0.999 확인','sram',25)
    elif v=='search_framework':
        flow(s,[('후보를 만드는 규칙','Panel · K chunk · loop 순서\ncache 위치를 조합','ir'),('검사 순서 선택','기본 space / 선택형 ReplayFunc\n동일한 NPU 후보 공간 사용','sram'),('실제 비용 확인','SRAM 배치 + 명령 생성\nDMA → words → SRAM','unit')],y=223,w=247,gap=42,h=131)
        s.table(['TVM 기능','현재 연결한 역할'],[
            ['Workload / JSONDatabase','계산 구조와 선택된 schedule trace 보관'],
            ['reuse / tune','기록 재생·검증 / 기존 winner와 새 후보 비교'],
            ['ReplayFunc','후보 선택을 반복; 학습형 비용 모델은 아님']], [305,521],y=400,row_h=60,size=18)
        s.tag(80,633,'같은 계산도 SRAM 용량·입출력 배치가 다르면 다시 확인','sram',796)
    elif v=='search_comparison':
        results={row['mode']:row for row in j('evidence/current/llama_schedule_refinement.json')['summary']}
        a,b=results['space'],results['meta_schedule']
        stat=lambda row,key:f"{row['statistics'][key]:,}"
        s.table(['512-token down / E9','Space','ReplayFunc'],[
            ['추가 합법 후보',f"{a['after_legal']}개",f"{b['after_legal']}개"],
            ['커널 탐색 시간',f"{a['tune_seconds']:.3f} s",f"{b['tune_seconds']:.3f} s"],
            ['최종 DMA bytes',stat(a,'dma_bytes'),stat(b,'dma_bytes')],
            ['최종 ISA words',stat(a,'words'),stat(b,'words')],
            ['최종 SRAM peak',stat(a,'sram_peak_bytes')+' B',stat(b,'sram_peak_bytes')+' B']], [300,263,263],y=222,row_h=55,size=18)
        s.card(80,555,798,84,'동일 binary · C-model 출력 검사 통과','실제 layer-0 weight + 합성 activation / 전체 Llama 출력 검증은 아님','sram',23)
    elif v=='full_search_pending':
        s.table(['256-token prefill','Generic off','Vectors','Matmul'],[
            ['Attention off','기준 구성','영역 10종','영역 11종'],
            ['Attention score','1개 구성','영역 9종','영역 10종'],
            ['Attention full','1개 구성','영역 9종','영역 10종']], [250,180,198,198],y=224,row_h=64,size=18)
        s.card(80,487,388,108,'전체 graph 구성 9개','같은 IR의 GQA·producer 조합은 제거\n영역 수는 종류이며 최종 채택 수가 아님','ir',24)
        s.card(487,487,391,108,'메모리 증가로 중단','기준 1개 완료 / 두 번째 구성에서 중단\n병목 수정 · 전체 출력 재검증은 남음','dram',22)
        s.tag(80,632,'28 layers 전체 · prefill만 · HF logits와 layer별 K/V 비교 예정','sram',796)
    elif v=='full_search':
        r=j('evidence/current/llama_prefill256_full_search.json')
        a=r['selection']['candidates'][0]['costs']['prefill_cache']
        b=r['entries']['prefill_cache']['statistics']
        s.table(['전체 prefill 지표','분리 실행 기준','선택된 구성'],[
            ['DMA bytes',f"{a['dma_bytes']:,}",f"{b['dma_bytes']:,}"],
            ['ISA words',f"{a['words']:,}",f"{b['words']:,}"],
            ['SRAM peak bytes',f"{a['sram_peak_bytes']:,}",f"{b['sram_peak_bytes']:,}"],
            ['Graph 시간 제한','없음 (None)','시간 종료 없음']], [230,298,298],y=225,row_h=66,size=18)
        ends=[e for e in r['candidate_events'] if e['event']=='candidate_end']
        s.card(80,525,798,82,f"28 layers · 256 tokens · graph {len(ends)}개 처리",
            'Matmul 목표 16개 / 최대 128회 · 분리 쪽도 SRAM 전달 적용','sram',22)
        s.text(84,625,f"HF 마지막 logits cosine: {r['validation']['cosine_vs_hf']:.9f} · token 일치",18,'teal',bold=True,w=790)
    elif v=='36':
        s.table(['검토 수준','대표 항목','먼저 확보할 근거'],[
            ['장치 계약 확인','Vector 256 / 긴 VLEN\nCycle와 실제 실행기','RTL 동작 / tail / buffer\n대표 연산 timing'],
            ['제한된 효과 비교','작은 fusion / KV 직접 저장\nCapacity bucket','최종 words·DMA·space\n구현과 ABI 변경 범위'],
            ['별도 큰 설계','Online softmax / tile streaming\nEviction / 정밀도 탐색','실제 손실이 큰 workload\n장치 상태와 검증 기준']], [169,341,316],y=227,row_h=108,size=18)
    elif v=='37':
        s.table(['단계','P = X × B를 따라가기','확인할 결과'],[
            ['모델 / Relax','P를 만들고 add/SiLU가 소비','Authored / graph prefix IR'],
            ['Semantic TIR','C32 누산 → P를 FP16로 저장','Reduction과 cast 위치'],
            ['Schedule','K 두 조각 / 입력 panel 재사용','Cache 위치와 intrinsic'],
            ['Memory','P의 수명 / SRAM 보관 / 주소','실제 선택과 byte offset'],
            ['ISA / runtime','LOAD·M_MUL·SAVE / 출력 해석','Program / plan / 실행 검증']], [166,350,310],y=224,row_h=66,size=18)
    elif v=='38':
        s.table(['표준 API','입력 → 결과','그 자체로 결정하지 않는 것'],[
            ['LegalizeOps / TE','고수준 연산 → 원소 수식','NPU의 허용 오차 정책'],
            ['FuseOps / FuseTIR','Graph 그룹 → 하나의 PrimFunc','실제 장치에서의 이득'],
            ['split / reorder','Loop 분할 / 의존성 내 순서 변경','최소 DMA 후보'],
            ['cache / compute_at','복사 stage / loop 위치 변경','공간에 맞는 탐색 정책'],
            ['tensorize','계산 조각 → 등록된 intrinsic','임의 연산의 ISA 구현'],
            ['Memory passes','수명·충돌 → storage 배치','호출 간 보관과 persistent KV']], [226,293,307],y=222,row_h=58,size=16)
    elif v=='39':
        code(s,78,219,803,99,'Relax: 함수 호출과 결과 shape','p = R.call_tir(matmul, (x, B), out_sinfo=R.Tensor((64,64), "float16"))','ir',16)
        code(s,78,352,803,113,'Semantic TIR: 원소의 정의','C32[i,j] += fp32(X[i,k]) * fp32(B[k,j])\nP[i,j] = fp16(C32[i,j])','unit',19)
        code(s,78,498,803,119,'Scheduled TIR: 복사와 tiled 실행 (간략한 의사 표현)','copy X/B panel to SRAM → for k_tile in range(2): npu_gemm_acc(...)\n전체 K 완료 → FP16 SAVE → add / SiLU 소비','sram',17)
        s.arrow(478,318,478,352);s.arrow(478,465,478,498)
    elif v=='40':
        s.table(['항목','Main ver.08','현재 v09'],[
            ['Lowering','Relax 직접 emitter','TIR + Schedule + 공통 codegen'],
            ['메모리','Flat G-buffer','DRAM ↔ SRAM ↔ unit'],
            ['Layer','Host가 같은 layer program 반복','전체 entry stream / 판단 cache'],
            ['Decode','Context 길이별 program','Capacity + mask / bucket'],
            ['KV','Host concat','별도 출력 + NPU append tail']], [140,302,384],y=228,row_h=67,size=18)
    elif v=='41':
        for yy,title,parts in [(251,'W8A16',[('INT8 W + scale','DRAM 저장','dram'),('FP16 W 복원','SRAM dequant','sram'),('FP16 GEMM','FP32 accumulator','unit')]),(452,'W8A8',[('INT8 W / 양자화 X','Per-channel / per-row','dram'),('INT8 dot64','INT32 부분합','unit'),('Scale 후 누산','FP32 → FP16 저장','sram')])]:
            s.text(79,yy-39,title,25,'teal',bold=True,w=799);flow(s,parts,y=yy,x=77,w=244,gap=41,h=120)
    elif v=='42':
        s.card(79,218,388,98,'전체 X: 18 MiB','64행 panel 384 KiB × 48개','sram',24)
        s.card(495,218,388,98,'전체 W: 48 MiB','256열 panel 1.5 MiB × 32개','ir',24)
        s.table(['바깥 loop','매 panel마다 다시 읽는 것','입력 전송 산술'],[
            ['X panel 48개','각 X panel에서 W 전체 순회','18 + 48×48 = 2322 MiB'],
            ['W panel 32개','각 W panel에서 X 전체 순회','48 + 18×32 = 624 MiB']], [179,315,332],y=365,row_h=81,size=18)
        s.tag(81,599,'출력·scratch 비용 제외 / 설명용 비교 / 실제 장문 compile 아님','dram',799)
    elif v=='43':
        flow(s,[('발표 슬라이드','그림 + 본문 + 예제 해석\n발표자 노트에 추가 설명','ir'),('주제별 문서','13개 장 / TVM API\n실행 계약 / 기존 main','sram'),('보존된 근거','실제 IR·JSON·명령\nSource hash와 실행 범위','dram')],y=227,w=247,gap=42,h=145)
        s.table(['파일','역할'],[
            ['slide_content.json','화면 설명·노트·근거의 편집 원본'],
            ['build_deck.py','편집 가능한 도형과 PPTX/PDF 재생성'],
            ['sources.json / validation.json','슬라이드별 출처 hash / 제작 검사 범위']], [371,455],y=429,row_h=59,size=18)
    else:raise ValueError(f'Unknown diagram {v}')

def make_slides():
    S.clear();items=json.loads((HERE/'slide_content.json').read_text())
    for n,item in enumerate(items,1):
        s=Slide(item,n,len(items));draw(s)


def emit(out,fontdir):
    out.mkdir(parents=True,exist_ok=True)
    for name,weight in [('NPU','Regular'),('NPUB','Bold')]:
        pdfmetrics.registerFont(TTFont(name,str(fontdir/f'NanumGothic-{weight}.ttf')))
    issues=[]
    for index,s in enumerate(S,1):
        for op in s.ops:
            if op['kind']=='text':
                font='NPUB' if op['bold'] else 'NPU'
                width=pdfmetrics.stringWidth(op['text'],font,op['size'])
                missing=[c for c in op['text'] if ord(c) not in pdfmetrics.getFont(font).face.charToGlyph]
                if width>op['w']+.2 or missing:
                    issues.append(dict(slide=index,text=op['text'],width=width,allowed=op['w'],missing=missing))
            if op['kind'] in ('text','box','image'):
                if op['x']<0 or op['y']<0 or op['x']+op['w']>W+.1 or op['y']+op['h']>H+.1:
                    issues.append(dict(slide=index,bounds=op))
    if issues:
        (out/'layout_issues.json').write_text(json.dumps(issues,ensure_ascii=False,indent=2)+'\n')
        raise ValueError(f'{len(issues)} layout/font issues; see layout_issues.json')
    (out/'layout_issues.json').unlink(missing_ok=True)
    deck=Presentation();deck.slide_width=Pt(W);deck.slide_height=Pt(H)
    deck.core_properties.title='TVM 기반 LLM NPU compiler: 구조와 동작 — 상세 설명판'
    deck.core_properties.subject='현재 v09 FP16 backend의 pipeline, 메모리 계획, 실행과 검증'
    deck.core_properties.author='NPU compiler project'
    pdf=canvas.Canvas(str(out/'npu_compiler_explained.pdf'),pagesize=(W,H))
    pdf.setTitle(deck.core_properties.title);pdf.setAuthor('NPU compiler project')
    def rgb(v):return RGBColor.from_string(v)
    def fill(v):pdf.setFillColorRGB(*[int(v[k:k+2],16)/255 for k in (0,2,4)])
    def stroke(v):pdf.setStrokeColorRGB(*[int(v[k:k+2],16)/255 for k in (0,2,4)])
    for s in S:
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text=s.title+'\n\n'+s.notes+'\n\n슬라이드 본문:\n'+'\n\n'.join(t['title']+'\n'+t['body'] for t in s.explain)+'\n\n예제와 해석: '+s.caption+'\n\n근거 파일 (docs 기준):\n'+'\n'.join(s.sources)
        for op in s.ops:
            kind=op['kind']
            if kind=='box':
                shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if op['radius'] else MSO_SHAPE.RECTANGLE,
                    Pt(op['x']),Pt(op['y']),Pt(op['w']),Pt(op['h']))
                if op['radius']:shape.adjustments[0]=min(.12,op['radius']/min(op['w'],op['h']))
                shape.fill.solid();shape.fill.fore_color.rgb=rgb(op['fill'])
                if op['stroke']:shape.line.color.rgb=rgb(op['stroke']);shape.line.width=Pt(1)
                else:shape.line.fill.background()
                fill(op['fill'])
                if op['stroke']:stroke(op['stroke'])
                pdf.roundRect(op['x'],H-op['y']-op['h'],op['w'],op['h'],op['radius'],fill=1,stroke=int(bool(op['stroke'])))
            elif kind=='text':
                box=slide.shapes.add_textbox(Pt(op['x']),Pt(op['y']),Pt(op['w']),Pt(op['h']))
                tf=box.text_frame;tf.clear();tf.word_wrap=False
                tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0
                tf.vertical_anchor=MSO_ANCHOR.TOP
                p=tf.paragraphs[0];p.space_before=p.space_after=Pt(0)
                p.line_spacing=Pt(op['size']*1.38)
                from pptx.enum.text import PP_ALIGN
                p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[op['align']]
                run=p.add_run();run.text=op['text'];run.font.name='NanumGothic';run.font.size=Pt(op['size'])
                run.font.bold=op['bold'];run.font.color.rgb=rgb(op['color'])
                for tag in ('a:ea','a:cs'):
                    el=OxmlElement(tag);el.set('typeface','NanumGothic');run._r.get_or_add_rPr().append(el)
                font='NPUB' if op['bold'] else 'NPU';pdf.setFont(font,op['size']);fill(op['color'])
                tx=op['x']
                if op['align']=='center':tx+= (op['w']-pdfmetrics.stringWidth(op['text'],font,op['size']))/2
                elif op['align']=='right':tx+=op['w']-pdfmetrics.stringWidth(op['text'],font,op['size'])
                pdf.drawString(tx,H-op['y']-op['size']*.94,op['text'])
            elif kind=='line':
                line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(op['x1']),Pt(op['y1']),Pt(op['x2']),Pt(op['y2']))
                line.line.color.rgb=rgb(op['color']);line.line.width=Pt(op['width'])
                stroke(op['color']);pdf.setLineWidth(op['width']);pdf.line(op['x1'],H-op['y1'],op['x2'],H-op['y2'])
            elif kind=='poly':
                points=op['points'];builder=slide.shapes.build_freeform(points[0][0],points[0][1],scale=Pt(1))
                builder.add_line_segments(points[1:],close=True);shape=builder.convert_to_shape()
                shape.fill.solid();shape.fill.fore_color.rgb=rgb(op['fill']);shape.line.fill.background()
                p=pdf.beginPath();p.moveTo(points[0][0],H-points[0][1])
                for x,y in points[1:]:p.lineTo(x,H-y)
                p.close();fill(op['fill']);pdf.drawPath(p,fill=1,stroke=0)
            elif kind=='image':
                slide.shapes.add_picture(op['path'],Pt(op['x']),Pt(op['y']),Pt(op['w']),Pt(op['h']))
                pdf.drawImage(op['path'],op['x'],H-op['y']-op['h'],op['w'],op['h'])
        pdf.showPage()
    pdf.save();deck.save(out/'npu_compiler_explained.pptx')
    notes=['# NPU compiler 발표자 노트 — 상세 설명판\n','본문 45장과 부록 6장. 본문 약 60~80분, 필요에 따라 두 회차로 나눈다.\n']
    for n,s in enumerate(S,1):
        notes += [f'## {n:02d}. {s.title}\n',s.notes+'\n', '\n'.join('**'+t['title']+'**\n\n'+t['body']+'\n' for t in s.explain), '**예제와 해석**\n\n'+s.caption+'\n', '근거: '+', '.join(f'[{p}](../{p})' for p in s.sources)+'\n']
    (out/'speaker_notes.md').write_text('\n'.join(notes)+'\n')
    paths=sorted({p for s in S for p in s.sources})
    sources=[]
    for p in paths:
        file=DOCS/p
        if not file.is_file():raise ValueError(f'Missing source {p}')
        sources.append(dict(path=p,sha256=hashlib.sha256(file.read_bytes()).hexdigest()))
    (out/'sources.json').write_text(json.dumps(dict(basis='Documentation snapshot, not a new compiler run',
        files=sources,slides=[dict(number=i,title=s.title,section=s.section,visual=s.visual,appendix=s.appendix,sources=s.sources) for i,s in enumerate(S,1)]),ensure_ascii=False,indent=2)+'\n')
    # Round trip through the PPTX parser verifies package relationships/notes.
    parsed=Presentation(out/'npu_compiler_explained.pptx')
    assert len(parsed.slides)==len(S)
    assert all(slide.notes_slide.notes_text_frame.text.strip() for slide in parsed.slides)
    check=dict(status='passed',slides=len(S),main_slides=sum(not s.appendix for s in S),
        appendix_slides=sum(s.appendix for s in S),sources=len(sources),
        checks=['text widths and shape bounds','Korean font coverage','PPTX reopen and speaker notes','source paths and hashes'],
        native_powerpoint_render_checked=False,pdf_method='Same scene geometry rendered with ReportLab; not a PowerPoint export',
        compiler_or_model_executed=False,
        edition='detailed', slide_size_points=[W,H],
        diagrams='Editable PowerPoint shapes; quantitative charts read from retained JSON',
        visible_text_characters=[sum(len(o['text']) for o in s.ops if o['kind']=='text') for s in S],
        presentation_source_sha256={name:hashlib.sha256((HERE/name).read_bytes()).hexdigest() for name in ['build_deck.py','slide_content.json']})
    try:
        import fitz
    except ImportError:
        check['pdf_pages_rendered']=None
    else:
        document=fitz.open(out/'npu_compiler_explained.pdf')
        assert len(document)==len(S)
        for page in document:
            page.get_pixmap(matrix=fitz.Matrix(.25,.25),alpha=False)
            assert page.get_text().strip() and '\ufffd' not in page.get_text()
        check['pdf_pages_rendered']=len(document)
        check['checks'].append('PDF page rasterization and text extraction')
        document.close()
    (out/'validation.json').write_text(json.dumps(check,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(check,ensure_ascii=False))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out',type=Path,default=HERE)
    parser.add_argument('--font-dir',type=Path,default=SUPPORT/'fonts')
    args=parser.parse_args()
    for name,weight in [('NPU','Regular'),('NPUB','Bold')]:
        pdfmetrics.registerFont(TTFont(name,str(args.font_dir/f'NanumGothic-{weight}.ttf')))
    make_slides();emit(args.out,args.font_dir)


if __name__=='__main__':main()
